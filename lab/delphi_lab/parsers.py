"""Bounded, read-only document adapters. Locators never claim Word page numbers.

Offsets in a locator use Python character offsets [start:end] into the original
paragraph, Markdown line or extracted PDF line, before whitespace normalization.
The immutable original file remains the authority for the container text.
"""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path, PurePosixPath
import re
from typing import Iterator
from zipfile import BadZipFile, ZipFile

from .models import Document, Side, SourceBlock

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
MAX_ZIP_ENTRIES = 2048
MAX_BLOCKS = 10000
MAX_TEXT_CHARS = 2_000_000
MAX_SEGMENT_CHARS = 8000
MAX_PDF_PAGES = 200
MAX_SHEETS = 20
MAX_ROWS = 10000
MAX_COLUMNS = 64


class ParseError(ValueError):
    """A format, safety or size failure; callers must not silently omit the file."""


@dataclass
class _RawBlock:
    text: str
    locator: str
    heading: bool = False
    toc: bool = False
    clause_numbers: bool = True


def normalize_text(text: str) -> str:
    """Normalize only a search copy. Never use this result as an original quote."""
    text = re.sub(r"\\([\\`*_{}\[\]()#+.!>\-])", r"\1", text)
    text = re.sub(r"^\s{0,3}#{1,6}\s+", "", text)
    text = text.replace("**", "").replace("__", "")
    return re.sub(r"\s+", " ", text).strip()


def _validate_zip(path: Path, required: str) -> list[str]:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ParseError("ZIP_ENTRY_LIMIT: too many archive entries")
            total = 0
            names = set()
            for item in entries:
                name = PurePosixPath(item.filename.replace("\\", "/"))
                if name.is_absolute() or ".." in name.parts:
                    raise ParseError("ZIP_INVALID_PATH: unsafe member name")
                if item.filename in names:
                    raise ParseError("ZIP_DUPLICATE_ENTRY: ambiguous archive member")
                names.add(item.filename)
                if item.flag_bits & 1:
                    raise ParseError("ZIP_ENCRYPTED: encrypted input is unsupported")
                total += item.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise ParseError("ZIP_EXPANSION_LIMIT: expanded archive is too large")
                if item.file_size > 1024 * 1024 and item.file_size / max(item.compress_size, 1) > 200:
                    raise ParseError("ZIP_RATIO_LIMIT: suspicious compression ratio")
            if required not in names or "[Content_Types].xml" not in names:
                raise ParseError("FORMAT_MISMATCH: archive does not match its extension")
            # These formats have no reason to contain DTDs. Inspect bounded XML
            # before handing it to a third-party parser; do not extract members.
            for item in entries:
                if item.filename.endswith((".xml", ".rels")):
                    xml = archive.read(item).upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise ParseError("XML_DTD_UNSUPPORTED: DTD/entity declarations are disallowed")
            return sorted(names)
    except BadZipFile as exc:
        raise ParseError("FORMAT_MISMATCH: invalid Office ZIP archive") from exc


def _docx(path: Path, names: list[str], warnings: list[str]) -> Iterator[_RawBlock]:
    from docx import Document as WordDocument
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    doc = WordDocument(path)
    body = doc.element.body
    if body.xpath(".//w:ins | .//w:del"):
        warnings.append("TRACKED_CHANGES_UNSUPPORTED: Word revisions require manual acceptance; extraction is incomplete.")
    if body.xpath(".//w:drawing | .//w:pict | .//w:txbxContent"):
        warnings.append("VISUAL_CONTENT_UNREAD: images, diagrams and text boxes are not interpreted.")
    if any(n.startswith("word/embeddings/") for n in names):
        warnings.append("EMBEDDED_FILES_UNREAD: embedded files are not recursively parsed.")
    ancillary = [n for n in names if re.match(r"word/(header|footer|footnotes|endnotes)\d*\.xml$", n)]
    if ancillary:
        from xml.etree import ElementTree
        with ZipFile(path) as archive:
            if any(any((node.text or "").strip() for node in ElementTree.fromstring(archive.read(n)).iter(qn("w:t"))) for n in ancillary):
                warnings.append("ANCILLARY_TEXT_UNREAD: headers, footers and notes contain text outside this adapter's body extraction.")
    warnings.append("DOCX_LOCATORS: paragraph/table positions and character offsets are used; Word pages are not inferred.")

    def walk(container, prefix: str = "") -> Iterator[_RawBlock]:
        paragraph_number = 0
        table_number = 0
        for child in container:
            if child.tag == qn("w:p"):
                paragraph_number += 1
                paragraph = Paragraph(child, doc)
                text = paragraph.text
                style = paragraph.style.name.lower() if paragraph.style is not None else ""
                if child.xpath("./w:pPr/w:numPr"):
                    warnings.append("AUTO_NUMBERING_UNRESOLVED: automatic list labels are not invented; use block locators.")
                yield _RawBlock(text, f"{prefix}paragraph {paragraph_number}", style.startswith("heading"), style.startswith("toc"))
            elif child.tag == qn("w:tbl"):
                table_number += 1
                for row_number, row in enumerate(child.findall(qn("w:tr")), 1):
                    for cell_number, cell in enumerate(row.findall(qn("w:tc")), 1):
                        yield from walk(cell, f"{prefix}table {table_number}, row {row_number}, cell {cell_number}, ")
            elif child.tag not in {qn("w:sectPr"), qn("w:tcPr")}:
                if child.xpath(".//w:t"):
                    warnings.append("UNSUPPORTED_WORD_CONTAINER: a structured content container has unextracted text.")

    yield from walk(body)


def _markdown(path: Path, warnings: list[str]) -> Iterator[_RawBlock]:
    try:
        text = path.read_bytes().decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ParseError("TEXT_ENCODING: Markdown must be UTF-8") from exc
    if "\x00" in text:
        raise ParseError("FORMAT_MISMATCH: binary bytes in Markdown")
    if re.search(r"!\[[^\]]*\]\(", text):
        warnings.append("MARKDOWN_IMAGES_UNREAD: linked images are not retrieved or interpreted.")
    # A line is a stable, exact quoting unit; continuation lines inherit context.
    for number, line in enumerate(text.splitlines(), 1):
        heading = bool(re.match(r"^\s*#{1,6}\s", line) or re.match(r"^\s*\*\*.*\*\*\s*$", line))
        yield _RawBlock(line, f"line {number}", heading)


def _pdf(path: Path, warnings: list[str]) -> Iterator[_RawBlock]:
    from pypdf import PdfReader

    with path.open("rb") as stream:
        reader = PdfReader(stream, strict=True)
        if reader.is_encrypted:
            raise ParseError("PDF_ENCRYPTED: password-protected PDF is unsupported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ParseError(f"PDF_PAGE_LIMIT: maximum is {MAX_PDF_PAGES} pages")
        warnings.append("PDF_READING_ORDER: extracted text order requires review for columns/tables; no OCR is performed.")
        for page_number, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                warnings.append(f"PDF_PAGE_UNREAD: page {page_number}, {type(exc).__name__}.")
                continue
            if not text.strip():
                warnings.append(f"PDF_PAGE_UNREAD: page {page_number} has no extractable text; OCR is unavailable.")
            if len(text) > MAX_TEXT_CHARS:
                raise ParseError("TEXT_LIMIT: oversized PDF page")
            for line_number, line in enumerate(text.splitlines(), 1):
                yield _RawBlock(line, f"page {page_number}, extracted line {line_number}")


def _xlsx(path: Path, warnings: list[str]) -> Iterator[_RawBlock]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    headers = {
        "unit": {"подразделение", "department", "unit", "бөлімше", "бөлім"},
        "function": {"функция", "обязанность", "function", "duty", "функциясы", "міндет", "міндеті"},
        "reporting": {"подчинённость", "подчиненность", "reports to", "reporting", "бағыныстылық", "бағынады"},
    }
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ParseError(f"XLSX_SHEET_LIMIT: maximum is {MAX_SHEETS} sheets")
        warnings.append("XLSX_TABULAR_ONLY: cell text is preserved; only Unit / Function / Reporting tables have an expected schema.")
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ParseError(f"XLSX_DIMENSION_LIMIT: sheet {sheet.title!r} exceeds {MAX_ROWS} rows / {MAX_COLUMNS} columns")
            found_header = False
            for row in sheet.iter_rows():
                nonempty = [c for c in row if c.value is not None]
                if not nonempty:
                    continue
                is_header = not found_header
                if is_header:
                    values = {normalize_text(str(c.value)).casefold() for c in nonempty}
                    recognized = all(values & aliases for aliases in headers.values())
                    if not recognized:
                        warnings.append(f"XLSX_SCHEMA_UNRECOGNIZED: sheet {sheet.title!r}; explicit cell interpretation is needed.")
                    found_header = True
                for cell in nonempty:
                    if cell.data_type == "f":
                        warnings.append("XLSX_FORMULAS_NOT_EVALUATED: formula text is preserved, not executed.")
                # Tab-delimited cell values retain the entire row. The header is
                # its parent source, so unit/function/reporting columns remain
                # together. Empty cells retain an empty field; no header words
                # or derived business meaning are inserted into the quotation.
                row_text = "\t".join("" if c.value is None else str(c.value) for c in row)
                locator = f"sheet {sheet.title!r}, cells A{nonempty[0].row}:{get_column_letter(len(row))}{nonempty[0].row} (tab-delimited values)"
                yield _RawBlock(row_text, locator, heading=is_header, clause_numbers=False)
    finally:
        workbook.close()


_START_NUMBER = re.compile(r"^(\d{1,2}(?:\.\d{1,3}){0,3})\.(?=\s|[^\W\d_])", re.UNICODE)
_EMBEDDED_NUMBER = re.compile(r"(?<=\s)(\d{1,2}(?:\.\d{1,3}){0,3})\.(?=[A-ZА-ЯӘІҢҒҮҰҚӨҺ])")


def _segments(text: str, clause_numbers: bool = True) -> Iterator[tuple[int, int]]:
    """Recognize explicit glued clause boundaries, then bound long fragments."""
    boundaries = [0]
    for match in _EMBEDDED_NUMBER.finditer(text) if clause_numbers else ():
        # A clause immediately followed by its capitalized sentence is a common
        # Word-export defect. Decimal references followed by whitespace do not
        # match, preventing a reference such as "п. 5.8.1" becoming a clause.
        boundaries.append(match.start())
    boundaries.append(len(text))
    for start, end in zip(boundaries, boundaries[1:]):
        while end - start > MAX_SEGMENT_CHARS:
            boundary = text.rfind(" ", start + MAX_SEGMENT_CHARS // 2, start + MAX_SEGMENT_CHARS)
            if boundary <= start:
                boundary = start + MAX_SEGMENT_CHARS
            yield start, boundary
            start = boundary
        if end > start:
            yield start, end


def _language(text: str) -> str:
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return "und"
    kk = sum(c in "әіңғүұқөһӘІҢҒҮҰҚӨҺ" for c in letters)
    cyrillic = sum("А" <= c <= "я" or c in "ЁёәіңғүұқөһӘІҢҒҮҰҚӨҺ" for c in letters)
    latin = sum("a" <= c.lower() <= "z" for c in letters)
    if kk >= 3 and kk / len(letters) >= .005:
        return "kk"
    if cyrillic / len(letters) > .6:
        return "ru"
    if latin / len(letters) > .6:
        return "en"
    return "mixed"


def _build_blocks(raw_blocks: Iterator[_RawBlock], document_id: str, side: Side, warnings: list[str]) -> list[SourceBlock]:
    blocks: list[SourceBlock] = []
    by_number: dict[str, SourceBlock] = {}
    numbered: SourceBlock | None = None
    unnumbered_heading: SourceBlock | None = None
    section = None
    heading_anchor: str | None = None
    in_toc = False
    total_chars = 0
    for raw in raw_blocks:
        total_chars += len(raw.text)
        if total_chars > MAX_TEXT_CHARS:
            raise ParseError("TEXT_LIMIT: extracted document is too large")
        for start, end in _segments(raw.text, raw.clause_numbers):
            original = raw.text[start:end]
            normalized = normalize_text(original)
            if not normalized:
                continue
            if len(blocks) >= MAX_BLOCKS:
                raise ParseError("BLOCK_LIMIT: too many extracted blocks")
            lower = normalized.casefold()
            match = _START_NUMBER.match(normalized) if raw.clause_numbers else None
            clause_no = match.group(1) if match else None
            parts = clause_no.split(".") if clause_no else []
            toc_title = lower in {"оглавление", "содержание", "contents", "table of contents", "мазмұны"}
            appendices_title = lower in {"приложения", "appendices", "appendix", "қосымшалар"}
            # A TOC marker is not permission to discard every following block.
            # Only explicit TOC styles, leader/tab + page patterns, or a short
            # paginated heading repeating an already-seen body heading qualify.
            # The last form handles this pair's trailing TOC, whose Word styles
            # are Heading rather than TOC and whose Markdown lines are bold.
            leader_page = bool(re.search(r"(?:\.{2,}|…{2,}|[·•]{2,}|\t+)\s*\d+\s*$", original))
            repeated_paginated_heading = False
            # Custom Word styles (for example RegHeading1) need not inherit
            # Heading. A previously seen numbered top-level heading is also
            # eligible, but only with matching title and an explicit page suffix.
            if in_toc and (raw.heading or len(parts) == 1) and clause_no in by_number and len(normalized) <= 240 and re.search(r"\s\d+\s*$", normalized):
                previous = by_number[clause_no]
                previous_number = _START_NUMBER.match(previous.normalized_text)
                if previous_number and previous.kind == "heading":
                    body_title = previous.normalized_text[previous_number.end():].strip().casefold()
                    entry_title = re.sub(r"\s+\d+(?:\s+.*)?$", "", normalized[match.end():].strip()).casefold()
                    repeated_paginated_heading = bool(entry_title and body_title.startswith(entry_title))
            is_toc = toc_title or raw.toc or (in_toc and (leader_page or repeated_paginated_heading))
            if toc_title:
                in_toc = True
            elif appendices_title:
                in_toc = False
                is_toc = raw.toc
            elif in_toc and not is_toc:
                # Keep the ambiguous boundary and all subsequent body text.
                # A later explicit TOC style or marker can still identify itself.
                warnings.append(f"TOC_BOUNDARY_REVIEW: content at {raw.locator} is retained as body because TOC continuation is not explicit.")
                in_toc = False
            kind = "toc" if is_toc else "heading" if raw.heading or (parts and len(parts) == 1) else "text"
            if kind == "text" and normalized.endswith(":") and len(normalized) <= 300:
                kind = "heading"
            parent = None
            if kind != "toc":
                if clause_no:
                    for depth in range(len(parts) - 1, 0, -1):
                        candidate = by_number.get(".".join(parts[:depth]))
                        if candidate:
                            parent = candidate.id
                            break
                    if len(parts) == 1:
                        section = parts[0]
                        unnumbered_heading = None
                        heading_anchor = None
                    elif unnumbered_heading:
                        # An unnumbered role can introduce several numbered
                        # duties (5 -> Chief -> 5.1/5.2). A new explicit
                        # numbered heading starts its own scope. An inner list
                        # heading under 9.5 must never become parent of 9.6.
                        inside_anchor = bool(heading_anchor and clause_no.startswith(heading_anchor + "."))
                        remainder = normalized[match.end():].strip() if match else normalized
                        explicit_role = bool(re.search(
                            r"^(?:(?:главный|chief|audit|senior|бас)\s+)?(?:директор\w*|руководител\w*|начальник\w*|аудитор\w*|director|head|manager|auditor|басшы\w*|директор\w*)\b",
                            remainder, re.I)) and not re.search(
                                r"\b(?:обязан\w*|долж\w*|вправе|must|shall|may|міндет\w*)\b", remainder, re.I)
                        starts_scope = kind == "heading" and (raw.heading or explicit_role)
                        if starts_scope or not inside_anchor:
                            unnumbered_heading = None
                            heading_anchor = None
                        elif parent == unnumbered_heading.parent_id:
                            parent = unnumbered_heading.id
                elif kind == "heading":
                    parent = numbered.id if numbered else by_number[section].id if section in by_number else None
                elif numbered:
                    parent = numbered.id
                elif unnumbered_heading:
                    parent = unnumbered_heading.id
            block = SourceBlock(
                id=f"{document_id}:b{len(blocks) + 1:05d}", document_id=document_id,
                side=side, locator=f"{raw.locator}, chars [{start}:{end}]",
                clause_no=clause_no, parent_id=parent,
                original_text=original, normalized_text=normalized, kind=kind,
            )
            blocks.append(block)
            if kind != "toc":
                if clause_no:
                    by_number[clause_no] = block
                    numbered = block
                    remainder = normalized[match.end():] if match else normalized
                    if not any(c.isalpha() for c in remainder):
                        warnings.append(f"EMPTY_CLAUSE: {clause_no} at {block.locator}; missing content is not reconstructed.")
                elif kind == "heading":
                    unnumbered_heading = block
                    anchor = next((b for b in reversed(blocks[:-1]) if b.id == parent), None)
                    heading_anchor = anchor.clause_no if anchor else None
                    numbered = None
    return blocks


def _appendix_warnings(blocks: list[SourceBlock], warnings: list[str]) -> None:
    # An appendix title ending in an Office filename, without any following
    # substantive content, identifies a reference rather than supplied content.
    for index, block in enumerate(blocks):
        if re.match(r"^(приложение|appendix|қосымша)\s+\d", block.normalized_text, re.I):
            following = [b for b in blocks[index + 1:] if b.kind != "toc"]
            if not following or (re.search(r"\.(?:docx?|pdf|xlsx?)\b", block.normalized_text, re.I) and len(following) < 2):
                warnings.append(f"APPENDIX_CONTENT_UNAVAILABLE: {block.id}; an appendix is named but its body is not supplied here.")


def parse_document(path: Path, side: Side, document_id: str | None = None) -> Document:
    """Parse a single local file without writing or fetching external resources.

    A ParseError is raised for unsafe/invalid/oversized files. "limited" reports
    successfully extracted evidence with a known coverage gap, never a silent
    successful skip. The optional document_id should be unique within a run.
    """
    path = Path(path)
    if side not in {"before", "after"}:
        raise ParseError("INVALID_SIDE: expected before or after")
    if not path.is_file():
        raise ParseError("FILE_NOT_FOUND: input is not a regular file")
    size = path.stat().st_size
    if size <= 0 or size > MAX_FILE_BYTES:
        raise ParseError(f"FILE_SIZE: expected 1..{MAX_FILE_BYTES} bytes")
    extension = path.suffix.lower()
    if extension not in {".docx", ".md", ".pdf", ".xlsx"}:
        raise ParseError("UNSUPPORTED_FORMAT: expected DOCX, UTF-8 MD, text PDF or XLSX")
    digest = sha256(path.read_bytes()).hexdigest()
    document_id = document_id or f"{side}-{digest[:16]}"
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,100}", document_id):
        raise ParseError("INVALID_DOCUMENT_ID: use 1..100 letters, digits, underscores or hyphens")
    warnings: list[str] = []
    if extension in {".docx", ".xlsx"}:
        required = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
        names = _validate_zip(path, required)
        adapter = _docx(path, names, warnings) if extension == ".docx" else _xlsx(path, warnings)
    elif extension == ".pdf":
        with path.open("rb") as stream:
            if not stream.read(8).startswith(b"%PDF-"):
                raise ParseError("FORMAT_MISMATCH: missing PDF signature")
        adapter = _pdf(path, warnings)
    else:
        adapter = _markdown(path, warnings)
    try:
        blocks = _build_blocks(adapter, document_id, side, warnings)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"EXTRACTION_FAILED: {type(exc).__name__}; cannot safely parse {extension}") from exc
    if not blocks:
        warnings.append("NO_TEXT: no extractable source blocks")
    _appendix_warnings(blocks, warnings)
    warnings = list(dict.fromkeys(warnings))
    informational = ("DOCX_LOCATORS:", "XLSX_TABULAR_ONLY:")
    limited = any(not warning.startswith(informational) for warning in warnings)
    return Document(
        id=document_id, side=side, filename=path.name, sha256=digest,
        format=extension[1:], detected_language=_language(" ".join(b.normalized_text for b in blocks)),
        parse_status="limited" if limited else "ok", warnings=warnings, blocks=blocks,
    )
