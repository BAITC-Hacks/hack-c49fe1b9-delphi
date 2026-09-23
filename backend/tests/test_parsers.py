from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.parsers import ParseError, parse_document

ROOT = Path(__file__).resolve().parents[2]
SIZE_LIMIT = 10_000_000


def parse(name: str, content: bytes):
    return parse_document(name, content, max_uncompressed_bytes=SIZE_LIMIT)


def find_clause(parsed, number: str):
    return next(block for block in parsed.blocks if block.clause_no == number)


@pytest.fixture
def revision_8():
    path = next((ROOT / "docs/sources").glob("*редакция_8*.docx"))
    return parse(path.name, path.read_bytes())


@pytest.fixture(params=["docx", "md"])
def revision_9(request):
    directory = ROOT / ("docs/sources" if request.param == "docx" else "docs/hackaton/tracks")
    path = next(directory.glob(f"*редакция_9*.{request.param}"))
    return parse(path.name, path.read_bytes())


def test_revision_9_docx_preserves_source_text_and_matches_markdown():
    docx_path = next((ROOT / "docs/sources").glob("*редакция_9*.docx"))
    md_path = next((ROOT / "docs/hackaton/tracks").glob("*редакция_9*.md"))
    docx = parse(docx_path.name, docx_path.read_bytes())
    markdown = parse(md_path.name, md_path.read_bytes())
    docx_clauses = {
        block.clause_no: block.normalized_text for block in docx.blocks if block.clause_no
    }
    md_clauses = {
        block.clause_no: block.normalized_text for block in markdown.blocks if block.clause_no
    }
    assert len(docx_clauses) == 318
    assert sum(block.clause_no is not None for block in docx.blocks) == 318
    assert docx_clauses == md_clauses
    assert "missing_annex_content:1" in docx.warnings
    paragraphs = Document(docx_path).paragraphs
    for block in docx.blocks:
        locator = block.locator
        original = paragraphs[locator["paragraph"] - 1].text
        assert original[locator["start_offset"] : locator["end_offset"]] == block.original_text


def test_real_docx_splits_joined_clauses_with_source_offsets(revision_8):
    blocks = [find_clause(revision_8, number) for number in ("3.9", "3.10", "3.11")]
    assert all(block.locator["paragraph"] == 122 for block in blocks)
    assert blocks[0].locator["end_offset"] == blocks[1].locator["start_offset"]
    assert blocks[1].locator["end_offset"] == blocks[2].locator["start_offset"]
    assert blocks[1].original_text.startswith("3.10.Работники")
    assert all("page" not in block.locator for block in revision_8.blocks)
    assert len({block.key for block in revision_8.blocks}) == len(revision_8.blocks)


def test_real_docx_empty_clause_annex_and_toc_are_explicit(revision_8):
    assert find_clause(revision_8, "5.5.3").original_text == "5.5.3. ;"
    assert "empty_clause:5.5.3" in revision_8.warnings
    assert "missing_annex_content:1" in revision_8.warnings
    assert any(warning.startswith("table_of_contents_skipped:") for warning in revision_8.warnings)
    assert len([block for block in revision_8.blocks if block.clause_no == "1"]) == 1
    assert revision_8.detected_language == "ru"
    assert revision_8.revision_label == "8"


@pytest.mark.parametrize("number", ["10", "11", "12", "13", "14"])
def test_real_docx_recovers_inline_section_headings(revision_8, number):
    assert find_clause(revision_8, number).parent_key is None


def test_real_documents_preserve_transferred_function_parent_context(revision_8, revision_9):
    for parsed, clause_no, owner_no in [(revision_8, "5.4.4", "5.4"), (revision_9, "5.3.3", "5.3")]:
        clause = find_clause(parsed, clause_no)
        owner = find_clause(parsed, owner_no)
        assert clause.parent_key == owner.key
        children = [block for block in parsed.blocks if block.parent_key == clause.key]
        assert len(children) == 2
        assert {block.locator["list_marker"] for block in children} == {"а", "б"}
        assert any("недостаточным или дублирующим" in block.original_text for block in children)
        assert "Директор" in owner.original_text
    assert "может осуществляться" not in find_clause(revision_8, "9.15").original_text
    assert "может осуществляться" in find_clause(revision_9, "9.15").original_text
    assert revision_9.revision_label == "9"
    assert revision_9.detected_language == "ru"


def test_markdown_preserves_original_and_does_not_promote_references_to_clauses():
    text = "**5\\. Duties**\n5.1. See clause 3.2. for details.\n5.2. Records are required. 6.Controls\n"
    parsed = parse("synthetic.md", text.encode())
    assert [block.clause_no for block in parsed.blocks] == ["5", "5.1", "5.2", "6"]
    assert parsed.blocks[0].original_text == "**5\\. Duties**"
    assert parsed.blocks[1].locator["line"] == 2
    assert parsed.blocks[1].parent_key == parsed.blocks[0].key


def test_docx_keeps_tables_between_paragraphs_and_reports_skipped_content():
    document = Document()
    document.add_paragraph("1. First department")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Audit"
    table.cell(0, 1).text = "Reviews controls"
    document.add_paragraph("2. Second department")
    document.element.body.append(OxmlElement("w:altChunk"))
    output = BytesIO()
    document.save(output)
    parsed = parse("synthetic.docx", output.getvalue())
    assert [block.original_text for block in parsed.blocks] == [
        "1. First department",
        "Audit",
        "Reviews controls",
        "2. Second department",
    ]
    assert parsed.blocks[1].locator["table"] == 1
    assert "docx_body_element_not_extracted:altChunk" in parsed.warnings


def test_xlsx_retains_header_context_and_cell_addresses():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Departments"
    sheet.append(["Department", "Function", "Reports to"])
    sheet.append(["Audit", "Review controls", "Board"])
    sheet.append(["Risk", "=1+1", "Board"])
    output = BytesIO()
    workbook.save(output)
    parsed = parse("synthetic.xlsx", output.getvalue())
    row = parsed.blocks[1]
    assert row.original_text == "Audit\tReview controls\tBoard"
    assert row.locator["headers"] == ["Department", "Function", "Reports to"]
    assert row.locator["cells"]["B2"] == "Review controls"
    assert row.locator["range"] == "A2:C2"
    assert "xlsx_formulas_not_evaluated:Departments:3" in parsed.warnings


def test_pdf_uses_real_pages_and_flags_page_without_text():
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 260 Td (1. Department audits controls.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(width=300, height=300)
    output = BytesIO()
    writer.write(output)
    parsed = parse("synthetic.pdf", output.getvalue())
    assert parsed.blocks[0].locator["page"] == 1
    assert parsed.blocks[0].clause_no == "1"
    assert "pdf_page_without_text:2" in parsed.warnings


@pytest.mark.parametrize(
    ("filename", "content", "code"),
    [
        ("file.exe", b"some content", "unsupported_format"),
        ("file.md", b"", "empty_file"),
        ("file.md", b"\xff", "invalid_encoding"),
        ("file.md", b"binary\x00text", "invalid_format"),
        ("file.md", b" \n ", "no_text"),
        ("file.docx", b"not a zip", "corrupt_document"),
        ("file.pdf", b"not a pdf", "invalid_format"),
    ],
)
def test_rejects_unreadable_or_unsupported_files(filename, content, code):
    with pytest.raises(ParseError) as caught:
        parse(filename, content)
    assert caught.value.code == code
    assert caught.value.message


def test_rejects_zip_with_missing_office_parts():
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("random.txt", "Not an Office document")
    with pytest.raises(ParseError, match="required parts"):
        parse("synthetic.docx", output.getvalue())


def test_rejects_expanded_size_before_parsing_archive():
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name in ["[Content_Types].xml", "_rels/.rels", "word/document.xml"]:
            archive.writestr(name, "x" * 1000)
    with pytest.raises(ParseError) as caught:
        parse_document("synthetic.docx", output.getvalue(), max_uncompressed_bytes=100)
    assert caught.value.code == "expanded_size_limit"


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("1. Бөлім қызметті тексереді.", "kk"),
        ("1. The department reviews controls.", "en"),
        ("1. 123", None),
    ],
)
def test_detects_content_language_without_word_metadata(text, expected):
    assert parse("synthetic.md", text.encode()).detected_language == expected


def test_unnumbered_role_context_applies_to_its_duties_only():
    text = (
        "5. Responsibilities\nChief auditor:\n"
        "5.1. Reviews controls.\n5.2. Reports to the board.\n"
        "5.3. Director of risk:\n5.3.1. Monitors operational risks.\n"
        "6. Resources\n6.1. Funding is provided.\n"
    )
    parsed = parse("synthetic-roles.md", text.encode())
    role = next(block for block in parsed.blocks if block.original_text == "Chief auditor:")
    assert role.parent_key == find_clause(parsed, "5").key
    assert find_clause(parsed, "5.1").parent_key == role.key
    assert find_clause(parsed, "5.2").parent_key == role.key
    assert find_clause(parsed, "5.3").parent_key == find_clause(parsed, "5").key
    assert find_clause(parsed, "5.3.1").parent_key == find_clause(parsed, "5.3").key
    assert find_clause(parsed, "6.1").parent_key == find_clause(parsed, "6").key


def test_inner_unnumbered_list_does_not_capture_next_sibling_duty():
    parsed = parse(
        "synthetic-scopes.md",
        (
            b"9. Audits\n9.5. Audit records include:\nRequired documents:\n"
            b"a) Audit plan\n9.6. Director approves the report.\n"
        ),
    )
    heading = next(block for block in parsed.blocks if block.original_text == "Required documents:")
    item = next(block for block in parsed.blocks if block.original_text == "a) Audit plan")
    assert item.parent_key == heading.key
    assert find_clause(parsed, "9.6").parent_key == find_clause(parsed, "9").key


def test_toc_marker_does_not_discard_ambiguous_body_ending_in_number():
    parsed = parse(
        "synthetic-toc.md",
        (
            b"Contents\n1. Responsibilities .... 3\n"
            b"1. Department must review records within 30\n1.1. Reports are required.\n"
        ),
    )
    assert "within 30" in find_clause(parsed, "1").original_text
    assert find_clause(parsed, "1.1").parent_key == find_clause(parsed, "1").key
    assert any(warning.startswith("toc_boundary_requires_review:") for warning in parsed.warnings)


def test_explicit_word_toc_style_is_not_extracted_as_duty():
    from docx.enum.style import WD_STYLE_TYPE

    document = Document()
    document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("1. Review .... 2", style="TOC 1")
    document.add_paragraph("1. Review actual controls.")
    output = BytesIO()
    document.save(output)
    parsed = parse("synthetic-toc.docx", output.getvalue())
    assert [block.original_text for block in parsed.blocks] == ["1. Review actual controls."]


def test_spreadsheet_headers_are_source_context_and_cell_numbers_are_not_clauses():
    workbook = Workbook()
    for sheet in (workbook.active, workbook.create_sheet("Second")):
        sheet.append(["Department", "Function", "Reports to"])
        sheet.append(["1. Audit", "Review clause 5.1. and report.", "Board"])
    output = BytesIO()
    workbook.save(output)
    parsed = parse("synthetic-table.xlsx", output.getvalue())
    assert len(parsed.blocks) == 4
    assert all(block.clause_no is None for block in parsed.blocks)
    assert parsed.blocks[1].parent_key == parsed.blocks[0].key
    assert parsed.blocks[2].parent_key is None
    assert parsed.blocks[3].parent_key == parsed.blocks[2].key
    assert parsed.blocks[1].original_text == "1. Audit\tReview clause 5.1. and report.\tBoard"


def test_spreadsheet_unrecognized_schema_is_explicit():
    workbook = Workbook()
    workbook.active.append(["Name", "Value"])
    workbook.active.append(["Audit", "Unknown meaning"])
    output = BytesIO()
    workbook.save(output)
    parsed = parse("synthetic-unknown-table.xlsx", output.getvalue())
    assert any(warning.startswith("xlsx_schema_unrecognized:") for warning in parsed.warnings)


def test_long_clauses_are_bounded_without_changing_quote_offsets():
    text = "5.1. " + "Review controls. " * 1100
    parsed = parse("synthetic-long.md", text.encode())
    assert len(parsed.blocks) > 1
    assert "".join(block.original_text for block in parsed.blocks) == text
    assert all(len(block.original_text) <= 8000 for block in parsed.blocks)
    for block in parsed.blocks:
        assert (
            text[block.locator["start_offset"] : block.locator["end_offset"]] == block.original_text
        )
    assert all(block.parent_key == parsed.blocks[0].key for block in parsed.blocks[1:])


@pytest.mark.parametrize(
    "declaration",
    [
        b'<!DOCTYPE document [<!ENTITY x "expanded">]>',
        "<!DOCTYPE document>".encode("utf-16"),
    ],
)
def test_office_dtd_is_rejected_before_xml_parser(declaration):
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("_rels/.rels", "<Relationships />")
        archive.writestr("word/document.xml", declaration + b"<document />")
    with pytest.raises(ParseError) as caught:
        parse("synthetic-dtd.docx", output.getvalue())
    assert caught.value.code == "unsupported_xml"


def test_office_unsafe_member_path_is_rejected():
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        for name in ("[Content_Types].xml", "_rels/.rels", "word/document.xml", "../outside.xml"):
            archive.writestr(name, "<empty />")
    with pytest.raises(ParseError) as caught:
        parse("synthetic-path.docx", output.getvalue())
    assert caught.value.code == "invalid_format"


def test_spreadsheet_dimension_limit_is_not_silently_truncated():
    workbook = Workbook()
    workbook.active.cell(row=10001, column=1, value="Review controls")
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ParseError) as caught:
        parse("synthetic-large.xlsx", output.getvalue())
    assert caught.value.code == "xlsx_dimension_limit"


def test_pdf_page_limit_is_enforced_before_extraction():
    writer = PdfWriter()
    for _ in range(201):
        writer.add_blank_page(width=30, height=30)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ParseError) as caught:
        parse("synthetic-large.pdf", output.getvalue())
    assert caught.value.code == "pdf_page_limit"
