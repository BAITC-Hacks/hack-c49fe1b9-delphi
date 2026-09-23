from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .types import TextChunk


def read_docx(content: bytes) -> tuple[list[TextChunk], list[str]]:
    document = Document(BytesIO(content))
    chunks: list[TextChunk] = []
    warnings: list[str] = []
    paragraph_index = 0
    table_index = 0
    for body_index, child in enumerate(document.element.body.iterchildren(), 1):
        if child.tag == qn("w:p"):
            paragraph_index += 1
            paragraph = Paragraph(child, document)
            style = paragraph.style.name.casefold() if paragraph.style is not None else ""
            locator = {"paragraph": paragraph_index, "body_index": body_index}
            chunks.append(
                TextChunk(
                    f"paragraph:{paragraph_index}",
                    paragraph.text,
                    locator,
                    heading=style.startswith("heading"),
                    toc=style.startswith("toc"),
                )
            )
            if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
                warnings.append(f"automatic_numbering_not_resolved:paragraph:{paragraph_index}")
        elif child.tag == qn("w:tbl"):
            table_index += 1
            table = Table(child, document)
            visited = set()
            for row_index, row in enumerate(table.rows, 1):
                for column_index, cell in enumerate(row.cells, 1):
                    if cell._tc in visited:
                        continue
                    visited.add(cell._tc)
                    for cell_index, paragraph in enumerate(cell.paragraphs, 1):
                        key = f"table:{table_index}:row:{row_index}:cell:{column_index}:paragraph:{cell_index}"
                        locator = {
                            "body_index": body_index,
                            "table": table_index,
                            "row": row_index,
                            "column": column_index,
                            "paragraph": cell_index,
                        }
                        style = (
                            paragraph.style.name.casefold() if paragraph.style is not None else ""
                        )
                        chunks.append(
                            TextChunk(
                                key,
                                paragraph.text,
                                locator,
                                heading=style.startswith("heading"),
                                toc=style.startswith("toc"),
                            )
                        )
                        if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
                            warnings.append(f"automatic_numbering_not_resolved:{key}")
                    if cell.tables:
                        warnings.append(
                            f"nested_tables_not_extracted:table:{table_index}:row:{row_index}"
                        )
        elif child.tag != qn("w:sectPr"):
            warnings.append(f"docx_body_element_not_extracted:{child.tag.rsplit('}', 1)[-1]}")
    unsupported = {
        "w:ins": "docx_tracked_changes_not_resolved",
        "w:del": "docx_tracked_changes_not_resolved",
        "w:txbxContent": "docx_textboxes_not_extracted",
        "w:drawing": "docx_drawings_not_extracted",
        "w:pict": "docx_drawings_not_extracted",
        "w:object": "docx_embedded_objects_not_extracted",
        "w:footnoteReference": "docx_footnotes_not_extracted",
        "w:endnoteReference": "docx_endnotes_not_extracted",
    }
    for tag, warning in unsupported.items():
        if next(document.element.iter(qn(tag)), None) is not None:
            warnings.append(warning)
    for relation in document.part.rels.values():
        if relation.reltype.rsplit("/", 1)[-1] in {"header", "footer"}:
            part = relation.target_part.element
            if any(node.text and node.text.strip() for node in part.iter(qn("w:t"))):
                warnings.append("docx_headers_footers_not_extracted")
    return chunks, warnings
