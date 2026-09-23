from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .limits import MAX_COLUMNS, MAX_ROWS, MAX_SHEETS
from .types import ParseError, TextChunk


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def read_xlsx(content: bytes) -> tuple[list[TextChunk], list[str]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    chunks = []
    warnings = []
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ParseError(
                "xlsx_sheet_limit", f"Spreadsheets are limited to {MAX_SHEETS} sheets."
            )
        for sheet in workbook:
            if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ParseError(
                    "xlsx_dimension_limit", "The spreadsheet dimensions exceed parser limits."
                )
            sheet.reset_dimensions()
            headers = None
            header_row = None
            for number, row in enumerate(sheet.iter_rows(), 1):
                if number > MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise ParseError(
                        "xlsx_dimension_limit", "The spreadsheet dimensions exceed parser limits."
                    )
                values = [_text(cell.value) for cell in row]
                if not any(values):
                    continue
                while values and not values[-1]:
                    values.pop()
                if headers is None:
                    headers = values
                    header_row = number
                    aliases = (
                        {"подразделение", "department", "unit", "бөлімше", "бөлім"},
                        {
                            "функция",
                            "обязанность",
                            "function",
                            "duty",
                            "функциясы",
                            "міндет",
                            "міндеті",
                        },
                        {
                            "подчинённость",
                            "подчиненность",
                            "reports to",
                            "reporting",
                            "бағыныстылық",
                            "бағынады",
                        },
                    )
                    normalized = {value.strip().casefold() for value in headers}
                    if not all(normalized & group for group in aliases):
                        warnings.append(f"xlsx_schema_unrecognized:{sheet.title}")
                if any(cell.data_type == "f" for cell in row):
                    warnings.append(f"xlsx_formulas_not_evaluated:{sheet.title}:{number}")
                locator = {
                    "sheet": sheet.title,
                    "row": number,
                    "range": f"A{number}:{get_column_letter(len(values))}{number}",
                    "headers": headers,
                    "header_row": header_row,
                    "cells": {
                        f"{get_column_letter(index)}{number}": value
                        for index, value in enumerate(values, 1)
                    },
                }
                chunks.append(
                    TextChunk(
                        f"sheet:{sheet.title}:row:{number}",
                        "\t".join(values),
                        locator,
                        heading=number == header_row,
                        clause_numbers=False,
                        context_key=f"sheet:{sheet.title}:row:{header_row}"
                        if number != header_row
                        else None,
                    )
                )
    finally:
        workbook.close()
    return chunks, warnings
